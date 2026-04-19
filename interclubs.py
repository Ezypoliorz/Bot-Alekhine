import time
import pandas as pd
from playwright.sync_api import sync_playwright
from bs4 import BeautifulSoup
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def parseRankingTable(htmlContent) :
    soupObj = BeautifulSoup(htmlContent, "html.parser")
    rankingTable = soupObj.find("table", {"id" : "TableClassement"})
    
    if not rankingTable or "Caen Alekhine" not in rankingTable.get_text() :
        return None

    rows = rankingTable.find_all("tr")
    parsedData = []
    
    for row in rows[1 :] :
        columns = row.find_all("td")
        if len(columns) < 7 :
            continue
            
        parsedData.append({
            "Pos." : columns[0].get_text(strip = True),
            "Équipe" : columns[1].get_text(strip = True),
            "Pts." : columns[2].get_text(strip = True).replace("\xa0", "0"),
            "Jouées" : columns[3].get_text(strip = True).replace("\xa0", "0"),
            "Diff." : columns[4].get_text(strip = True),
            "Pro." : columns[5].get_text(strip = True),
            "Con." : columns[6].get_text(strip = True)
        })

    return pd.DataFrame(parsedData)

def parseMatchDetails(htmlContent) :
    soupObj = BeautifulSoup(htmlContent, "html.parser")
    allRows = soupObj.find_all("tr")
    
    matchDataFrames = []
    isTargetMatch = False
    currentMatchRows = []
    
    for row in allRows :
        rowId = row.get("id", "")
        
        if "RowMatchTitre" in rowId :
            if isTargetMatch and currentMatchRows :
                matchDataFrames.append(pd.DataFrame(currentMatchRows))
                currentMatchRows = []
            
            if "Caen Alekhine" in row.get_text() :
                isTargetMatch = True
                cells = row.find_all(["td", "th"])
                currentMatchRows.append([c.get_text(strip = True) for c in cells])
            else :
                isTargetMatch = False
        
        elif "RowMatchDetail" in rowId and isTargetMatch :
            cells = row.find_all(["td", "th"])
            currentMatchRows.append([c.get_text(strip = True) for c in cells])
            
    if isTargetMatch and currentMatchRows :
        matchDataFrames.append(pd.DataFrame(currentMatchRows))
        
    if matchDataFrames :
        finalList = []
        for df in matchDataFrames :
            finalList.append(df)
        return pd.concat(finalList, ignore_index = True)
        
    return None

def applyStyle(sheet, startRow, startCol, dataFrame, isTitle = False, hasHeader = True, isMainHeader = False, colSpan = None, boldFirstDataRow = False) :
    headerFill = PatternFill(start_color = "CFE2F3", end_color = "CFE2F3", fill_type = "solid")
    titleFill = PatternFill(start_color = "4A86E8", end_color = "4A86E8", fill_type = "solid")
    headerFont = Font(bold = True)
    whiteFont = Font(bold = True, color = "FFFFFF", size = 14)
    centerAlignment = Alignment(horizontal = "center", vertical = "center")
    thinSide = Side(style = "thin", color = "000000")
    blackBorder = Border(left = thinSide, right = thinSide, top = thinSide, bottom = thinSide)

    if isMainHeader :
        cell = sheet.cell(row = startRow, column = startCol)
        cell.font = whiteFont
        cell.fill = titleFill
        cell.alignment = centerAlignment
        useColSpan = colSpan if colSpan is not None else 1
        for rIdx in range(startRow, startRow + 2) :
            for cIdx in range(startCol, startCol + useColSpan) :
                sheet.cell(row = rIdx, column = cIdx).border = blackBorder
        return

    if isTitle :
        cell = sheet.cell(row = startRow, column = startCol)
        cell.font = headerFont
        cell.fill = headerFill
        cell.alignment = centerAlignment
        useColSpan = colSpan if colSpan is not None else 1
        for cIdx in range(useColSpan) :
            sheet.cell(row = startRow, column = startCol + cIdx).border = blackBorder
        return

    numRows = len(dataFrame)
    numCols = colSpan if colSpan is not None else len(dataFrame.columns)

    if hasHeader :
        for colIdx in range(numCols) :
            cell = sheet.cell(row = startRow, column = startCol + colIdx)
            cell.font = headerFont
            cell.fill = headerFill
            cell.alignment = centerAlignment
            cell.border = blackBorder
        dataStartRow = startRow + 1
    else :
        dataStartRow = startRow

    for rIdx in range(numRows) :
        for cIdx in range(numCols) :
            cell = sheet.cell(row = dataStartRow + rIdx, column = startCol + cIdx)
            cell.border = blackBorder
            cell.alignment = centerAlignment
            if boldFirstDataRow and rIdx == 0 :
                cell.font = headerFont

def autoAdjustColumns(sheet) :
    for col in sheet.columns :
        maxLength = 0
        columnLetter = col[0].column_letter
        for cell in col :
            try :
                if cell.value :
                    maxLength = max(maxLength, len(str(cell.value)))
            except : pass
        sheet.column_dimensions[columnLetter].width = max(maxLength + 3, 10)

def runSession() :
    with sync_playwright() as p :
        browser = p.chromium.launch(headless = False)
        context = browser.new_context()
        page = context.new_page()
        page.route("**/*.{png,jpg,jpeg,gif,css,woff,woff2}", lambda route : route.abort())
        
        fileName = "Interclubs adultes Caen Alekhine.xlsx"
        excelWriter = pd.ExcelWriter(fileName, engine = "openpyxl")
        foundAnyData = False
        
        print("Connexion au site FFE...")
        page.goto("https://www.echecs.asso.fr/Equipes.aspx")

        seasonIndex = 0
        while True :
            seasonSelector = f"id=ctl00_ContentPlaceHolderMain_RepeaterSaisons_ctl{seasonIndex :02d}_SaisonBase"
            if not page.query_selector(seasonSelector) :
                break
            
            fullYearText = page.inner_text(seasonSelector)
            shortYear = fullYearText.split("-")[0].strip()
            with page.expect_navigation() :
                page.click(seasonSelector)
            
            try :
                page.wait_for_selector("select[name='ctl00$ContentPlaceHolderMain$SelectCompetition']", timeout = 5000)
                options = page.query_selector_all("select[name='ctl00$ContentPlaceHolderMain$SelectCompetition'] option")
                targetValue = next((o.get_attribute("value") for o in options if "interclub" in o.inner_text().lower()), None)
                if not targetValue :
                    seasonIndex += 2
                    continue
                with page.expect_navigation() :
                    page.select_option("select[name='ctl00$ContentPlaceHolderMain$SelectCompetition']", value = targetValue)
            except Exception :
                seasonIndex += 2
                continue

            divOptions = page.query_selector_all("select[name='ctl00$ContentPlaceHolderMain$SelectDivision'] option")
            divisionList = [{"val" : o.get_attribute("value"), "name" : o.inner_text()} for o in divOptions if o.get_attribute("value") != "0"]

            for div in divisionList :
                with page.expect_navigation() :
                    page.select_option("select[name='ctl00$ContentPlaceHolderMain$SelectDivision']", value = div["val"])
                
                grpOptions = page.query_selector_all("select[name='ctl00$ContentPlaceHolderMain$SelectGroupe'] option")
                groupList = [{"val" : o.get_attribute("value"), "name" : o.inner_text()} for o in grpOptions if o.get_attribute("value") != "0"]

                for grp in groupList :
                    with page.expect_navigation() :
                        page.select_option("select[name='ctl00$ContentPlaceHolderMain$SelectGroupe']", value = grp["val"])
                    
                    dfRanking = parseRankingTable(page.content())
                    if dfRanking is not None :
                        sheetName = f"{div['name']} - {shortYear}"[:31].replace("/", "-")
                        pd.DataFrame().to_excel(excelWriter, sheet_name = sheetName)
                        workbook = excelWriter.book
                        currentSheet = workbook[sheetName]
                        
                        colStartLeft = 2
                        colStartRight = 6
                        colEnd = 8
                        
                        mainTitleText = f"{div['name']} - {grp['name']} ({shortYear})"
                        currentSheet.merge_cells(start_row = 2, start_column = colStartLeft, end_row = 3, end_column = colEnd)
                        applyStyle(currentSheet, startRow = 2, startCol = colStartLeft, dataFrame = pd.DataFrame(), isMainHeader = True, colSpan = 7)
                        currentSheet.cell(row = 2, column = colStartLeft).value = mainTitleText

                        rankingStartRow = 5
                        dfRanking.to_excel(excelWriter, sheet_name = sheetName, index = False, startrow = rankingStartRow - 1, startcol = colStartLeft - 1)
                        applyStyle(currentSheet, startRow = rankingStartRow, startCol = colStartLeft, dataFrame = dfRanking)
                        
                        roundsStartRow = rankingStartRow + len(dfRanking) + 2
                        leftRow = roundsStartRow
                        rightRow = roundsStartRow
                        useLeft = True
                        
                        roundLinks = page.query_selector_all("a[id*='LinkCmdRonde']")
                        for i in range(len(roundLinks)) :
                            try :
                                currentLinks = page.query_selector_all("a[id*='LinkCmdRonde']")
                                roundLabel = currentLinks[i].inner_text()
                                with page.expect_navigation(timeout = 120000) :
                                    currentLinks[i].click()

                                dfMatch = parseMatchDetails(page.content())
                                if dfMatch is not None :
                                    activeCol = colStartLeft if useLeft else colStartRight
                                    activeRow = leftRow if useLeft else rightRow
                                    
                                    currentSheet.merge_cells(start_row = activeRow, start_column = activeCol, end_row = activeRow, end_column = activeCol + 2)
                                    applyStyle(currentSheet, startRow = activeRow, startCol = activeCol, dataFrame = pd.DataFrame(), isTitle = True, colSpan = 3)
                                    currentSheet.cell(row = activeRow, column = activeCol).value = f"{roundLabel}"
                                    
                                    dfMatch.to_excel(excelWriter, sheet_name = sheetName, index = False, header = False, startrow = activeRow, startcol = activeCol - 1)
                                    applyStyle(currentSheet, startRow = activeRow + 1, startCol = activeCol, dataFrame = dfMatch, hasHeader = False, colSpan = 3, boldFirstDataRow = True)
                                    
                                    if useLeft :
                                        leftRow += len(dfMatch) + 3
                                    else :
                                        rightRow += len(dfMatch) + 3
                                    useLeft = not useLeft

                                page.go_back(timeout = 90000)
                                time.sleep(1)
                            except Exception :
                                continue
                        
                        autoAdjustColumns(currentSheet)
                        foundAnyData = True
            seasonIndex += 2

        if foundAnyData :
            excelWriter.close()
            print(f"\nFichier Excel genere.")
        browser.close()

if __name__ == "__main__" :
    runSession()